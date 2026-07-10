from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


INPUT = Path("/Users/a0000/Desktop/论文/GMCP-R_SCI中文初稿_内容修订版_2.6.docx")
OUTPUT = Path("/Users/a0000/Desktop/论文/GMCP-R_SCI中文初稿_内容修订版_2.7_按建议修订.docx")
ASSET_DIR = Path("/Users/a0000/Desktop/实验/gmcp_r/gmcpr_revision_assets")


def load_font(size: int, mono: bool = False):
    candidates = []
    if mono:
        candidates.extend(
            [
                "/System/Library/Fonts/Menlo.ttc",
                "/System/Library/Fonts/SFNSMono.ttf",
                "/Library/Fonts/Arial Unicode.ttf",
            ]
        )
    candidates.extend(
        [
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Light.ttc",
            "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
            "/Library/Fonts/Arial Unicode.ttf",
        ]
    )
    for candidate in candidates:
        try:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


def draw_arrow(draw, start, end, color=(54, 64, 78), width=4):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - 18 * direction, y2 - 10), (x2 - 18 * direction, y2 + 10)]
    else:
        direction = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 10, y2 - 18 * direction), (x2 + 10, y2 - 18 * direction)]
    draw.polygon(pts, fill=color)


def rounded_box(draw, xy, text, font, fill, outline=(80, 95, 115), text_fill=(24, 32, 44)):
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    lines = text.split("\n")
    line_heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total_h = sum(line_heights) + (len(lines) - 1) * 10
    y = y1 + (y2 - y1 - total_h) / 2
    for line, h in zip(lines, line_heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=text_fill)
        y += h + 10


def create_memory_chain_figure(path: Path):
    w, h = 1650, 980
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(44)
    node_font = load_font(38, mono=True)
    text_font = load_font(26, mono=True)
    small_font = load_font(24, mono=True)

    draw.text((70, 45), "GMCP-R Memory Chain State Evolution", font=title_font, fill=(20, 38, 62))
    draw.line((70, 105, 1580, 105), fill=(180, 190, 205), width=2)

    node_x = 250
    node_ys = [190, 385, 580, 775]
    labels = ["M0", "M1", "M2", "M3"]
    for y, label in zip(node_ys, labels):
        draw.ellipse((node_x - 58, y - 58, node_x + 58, y + 58), fill=(235, 243, 255), outline=(48, 91, 150), width=4)
        bbox = draw.textbbox((0, 0), label, font=node_font)
        draw.text((node_x - (bbox[2] - bbox[0]) / 2, y - (bbox[3] - bbox[1]) / 2 - 4), label, font=node_font, fill=(26, 58, 100))

    for i in range(3):
        draw_arrow(draw, (node_x, node_ys[i] + 62), (node_x, node_ys[i + 1] - 62), color=(48, 91, 150), width=5)

    steps = [
        ("msg1: seq=1, h1, prev_mem=M0", "M1 = H(M0 || sid || epoch || seq1 || h1 || sender)"),
        ("msg2: seq=2, h2, prev_mem=M1", "M2 = H(M1 || sid || epoch || seq2 || h2 || sender)"),
        ("msg3: seq=3, h3, prev_mem=M2", "M3 = H(M2 || sid || epoch || seq3 || h3 || sender)"),
    ]
    for y, (line1, line2) in zip([292, 487, 682], steps):
        rounded_box(
            draw,
            (420, y - 70, 1540, y + 70),
            f"{line1}\n{line2}",
            small_font,
            fill=(248, 251, 255),
            outline=(137, 162, 195),
        )
        draw.line((365, y, 420, y), fill=(137, 162, 195), width=3)

    note = "Receiver check:\npacket.prev_mem == local.last_mem; update local.last_mem = Mi"
    rounded_box(draw, (420, 815, 1540, 940), note, small_font, fill=(244, 250, 246), outline=(103, 151, 116))
    img.save(path)


def create_protocol_framework_figure(path: Path):
    w, h = 1650, 1040
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(44)
    head_font = load_font(34)
    text_font = load_font(25, mono=True)
    small_font = load_font(23, mono=True)

    draw.text((70, 45), "GMCP-R Protocol Framework and Recovery Path", font=title_font, fill=(20, 38, 62))
    draw.line((70, 105, 1580, 105), fill=(180, 190, 205), width=2)

    client_x, server_x = 250, 1220
    draw.text((client_x - 70, 135), "Client", font=head_font, fill=(24, 55, 92))
    draw.text((server_x - 70, 135), "Server", font=head_font, fill=(24, 55, 92))
    draw.line((client_x, 185, client_x, 940), fill=(205, 214, 226), width=4)
    draw.line((server_x, 185, server_x, 940), fill=(205, 214, 226), width=4)

    # DATA phase.
    y = 240
    rounded_box(draw, (70, y - 50, 430, y + 50), "DATA\nseq_i, payload_i,\nprev_mem", small_font, fill=(239, 246, 255), outline=(72, 117, 180))
    draw_arrow(draw, (430, y), (1040, y), color=(48, 91, 150), width=5)
    rounded_box(
        draw,
        (1040, y - 120, 1570, y + 120),
        "verify HMAC\nverify seq\nverify prev_mem == last_mem\nupdate last_seq, last_mem",
        small_font,
        fill=(248, 251, 255),
        outline=(137, 162, 195),
    )

    y = 440
    rounded_box(draw, (1030, y - 65, 1570, y + 65), "periodic checkpoint\nckpt_seq, ckpt_mem", small_font, fill=(244, 250, 246), outline=(103, 151, 116))
    draw_arrow(draw, (1040, y + 85), (430, y + 85), color=(72, 125, 96), width=5)
    rounded_box(draw, (70, y + 35, 430, y + 135), "ACK /\nMemoryTicket", small_font, fill=(244, 250, 246), outline=(103, 151, 116))

    # Gap.
    y = 610
    draw.rounded_rectangle((520, y - 40, 1110, y + 40), radius=20, fill=(255, 248, 235), outline=(202, 151, 65), width=3)
    gap = "disconnect / attack / state gap"
    bbox = draw.textbbox((0, 0), gap, font=text_font)
    draw.text((815 - (bbox[2] - bbox[0]) / 2, y - 18), gap, font=text_font, fill=(115, 80, 30))

    # Recovery phase.
    y = 760
    rounded_box(draw, (70, y - 70, 430, y + 70), "RECOVERY_REQUEST\n+\nMemoryTicket", small_font, fill=(255, 242, 242), outline=(176, 85, 85))
    draw_arrow(draw, (430, y), (1040, y), color=(176, 85, 85), width=5)
    rounded_box(
        draw,
        (1040, y - 125, 1570, y + 125),
        "verify ticket\ncheck nonce / expire\ncheck rollback / session\nrecover last_seq, last_mem",
        small_font,
        fill=(255, 248, 248),
        outline=(176, 85, 85),
    )

    y = 930
    draw_arrow(draw, (1040, y), (430, y), color=(48, 91, 150), width=5)
    rounded_box(draw, (70, y - 50, 430, y + 50), "RECOVERY_RESPONSE\nlast_seq, last_mem,\ncheckpoint", small_font, fill=(239, 246, 255), outline=(72, 117, 180))
    img.save(path)


def set_text(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    apply_document_font(run)
    return run


def apply_document_font(run):
    run.font.name = "Times New Roman"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "Songti SC")
    r_fonts.set(qn("w:cs"), "Times New Roman")


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph.__class__(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        apply_document_font(run)
    return new_para


def insert_paragraph_after_table(table, text="", style=None):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, table._parent)
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        apply_document_font(run)
    return new_para


def find_paragraph(doc, predicate):
    for paragraph in doc.paragraphs:
        if predicate(paragraph.text):
            return paragraph
    raise ValueError("target paragraph not found")


def set_cell_text(cell, text, bold=False, font_size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    apply_document_font(run)
    run.bold = bold
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_width(cell, width_in):
    width = Inches(width_in)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width.twips)))


def format_paragraph(paragraph, size=10.5, bold=False, align=None):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.bold = bold


def add_caption_after(paragraph, caption):
    cap = insert_paragraph_after(paragraph, caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        run.font.size = Pt(10)
    return cap


def add_figure_after(anchor, image_path: Path, caption: str):
    p = insert_paragraph_after(anchor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(6.3))
    return add_caption_after(p, caption)


def replace_table_data(table, data, widths):
    while len(table.rows) > len(data):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(data):
        table.add_row()
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_text(cell, data[row_idx][col_idx], bold=(row_idx == 0))
            set_cell_width(cell, widths[col_idx])
    set_repeat_header(table)


def set_repeat_header(table):
    if not table.rows:
        return
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def main():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    memory_fig = ASSET_DIR / "gmcpr_memory_chain.png"
    framework_fig = ASSET_DIR / "gmcpr_protocol_framework.png"
    create_memory_chain_figure(memory_fig)
    create_protocol_framework_figure(framework_fig)

    doc = Document(INPUT)
    for table in doc.tables:
        set_repeat_header(table)

    # Abstract and keywords.
    abstract = find_paragraph(doc, lambda t: t.startswith("连接恢复并不等同于历史状态恢复"))
    set_text(
        abstract,
        "连接恢复并不等同于历史状态恢复。针对断续通信中恢复后历史连续性难以验证的问题，本文提出 "
        "GMCP-R（Memory-Continuity-Aware Communication Protocol with Recovery）。GMCP-R 解决的是断连或攻击检测后"
        "“如何确认恢复后的通信历史仍然连续可信”的问题；其恢复目标不是单纯重新建立 TCP 连接，而是恢复由 last_seq、"
        "last_mem 与 Checkpoint 共同刻画的历史记忆状态。具体而言，last_seq 表示可接受的恢复点，last_mem 表示截至该恢复点"
        "处理后的链式历史状态，Checkpoint 提供最近可信快照以限制恢复重放范围。该协议结合 memory chain、MemoryTicket "
        "与 Checkpoint：memory chain 将消息内容、序列号、会话上下文和前序记忆状态绑定；MemoryTicket 以服务器认证票据"
        "保存 last_seq、last_mem 与 checkpoint 信息；Checkpoint 用于限制恢复时的历史重放范围。基于 Python TCP 原型的实验"
        "表明，在测试攻击场景下，GMCP-R 达到 100% 检测率、0% 误接受率和 0% 误拒绝率，五类无效 MemoryTicket 均被拒绝；"
        "本地纯计算吞吐量均值为 65,931 msg/s，1 至 20 个并发客户端均保持 100% 成功率，吞吐量峰值为 17,927 msg/s。"
        "弱网结果仅来自 preliminary code-level simulation，用于补充观察恢复行为，不作为核心结论。结果说明，GMCP-R 可在"
        "可接受开销下为断续通信提供历史连续性验证与安全恢复支持。"
    )
    keywords = find_paragraph(doc, lambda t: t.startswith("Keywords（关键词）"))
    set_text(keywords, "Keywords（关键词）：安全通信；历史连续性；状态恢复；断续通信；MemoryTicket；Checkpoint；攻击检测；通信协议")

    # Section 1.2 recovery-point definition.
    problem_para = find_paragraph(doc, lambda t: t.startswith("本文关注断续通信中的历史连续性恢复问题"))
    set_text(
        problem_para,
        "本文关注断续通信中的历史连续性恢复问题：当通信在第 i 条消息附近发生异常中断或攻击检测后，恢复机制不仅应恢复会话，"
        "还应验证恢复点之前的消息历史与双方维护的记忆状态一致。换言之，恢复后的 last_seq 与 last_mem 必须能解释为同一条"
        "历史链上的状态，而不是由攻击者诱导出的旧状态、跨会话状态或伪造状态。"
    )
    inserted = insert_paragraph_after(
        problem_para,
        "在本文中，恢复点由 last_seq 表示，恢复状态由 last_mem 表示。last_seq 表示通信双方已经接收并接受的最新消息序号，"
        "last_mem 表示截至该序号处理后的历史记忆状态；Checkpoint 则记录最近可信的 {ckpt_seq, ckpt_mem} 快照，用于缩短恢复时"
        "需要验证或重放的历史窗口。一次安全恢复并不是简单重连，而是要求恢复后的 last_seq 与 last_mem 能够对应"
        "同一条合法历史链。若攻击者试图使用旧票据、跨会话票据或伪造记忆状态恢复通信，则服务端应拒绝该恢复请求。",
    )
    format_paragraph(inserted)

    contribution = find_paragraph(doc, lambda t: t.startswith("1. 定义断续通信中的历史连续性恢复问题"))
    set_text(
        contribution,
        "1. 定义断续通信中的历史连续性恢复问题，区分连接恢复、消息完整性验证与历史状态恢复，并明确恢复点 last_seq 与恢复状态 last_mem 的语义。",
    )

    # Remove teacher-note text from headings.
    set_text(find_paragraph(doc, lambda t: t.startswith("4.2 Memory Chain")), "4.2 Memory Chain")
    set_text(find_paragraph(doc, lambda t: t.startswith("4.3 DATA 报文格式")), "4.3 DATA 报文格式")
    set_text(find_paragraph(doc, lambda t: t.startswith("5. Security Analysis")), "5. Security Analysis（安全分析）")

    # Renumber existing result figures before adding two protocol figures.
    fig_updates = {
        "图 1. 各协议攻击检测率与误接受率对比。": "图 3. 各协议攻击检测率与误接受率对比。",
        "图 2. 正常通信条件下各协议平均吞吐量对比。": "图 4. 正常通信条件下各协议平均吞吐量对比。",
        "图 3. 正常通信条件下各协议平均 RTT 对比。": "图 5. 正常通信条件下各协议平均 RTT 对比。",
        "图 4. 本地纯计算吞吐量随载荷大小变化。": "图 6. 本地纯计算吞吐量随载荷大小变化。",
        "图 5. 并发客户端数量与总吞吐量变化。": "图 7. 并发客户端数量与总吞吐量变化。",
        "图 6. 并发客户端数量与平均 RTT 变化。": "图 8. 并发客户端数量与平均 RTT 变化。",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in fig_updates:
            set_text(paragraph, fig_updates[text])

    # Protocol overview figure.
    overview = find_paragraph(doc, lambda t: t.startswith("GMCP-R 在正常通信阶段"))
    bridge = insert_paragraph_after(
        overview,
        "从整体流程看，GMCP-R 将正常 DATA 通信、周期性 Checkpoint、MemoryTicket 签发与异常后的恢复请求放在同一状态框架内。"
        "DATA 阶段持续推进 last_seq 与 last_mem，恢复阶段则以 MemoryTicket 和 Checkpoint 验证恢复点及恢复状态是否仍属于同一条合法历史链。",
    )
    format_paragraph(bridge)
    add_figure_after(bridge, framework_fig, "图 1. GMCP-R 协议整体框架与恢复路径。")

    # Memory chain figure.
    memory_para = find_paragraph(doc, lambda t: t.startswith("其中，sid 表示会话标识"))
    bridge2 = insert_paragraph_after(
        memory_para,
        "图 2 展示了 memory chain 的状态演进关系。每条消息均携带发送前的 prev_mem，接收端只有在 prev_mem 与本地 last_mem 一致时才接受该消息，并在认证通过后计算新的 last_mem。",
    )
    format_paragraph(bridge2)
    add_figure_after(bridge2, memory_fig, "图 2. GMCP-R 的 memory chain 状态演进过程。")

    # DATA packet table and optional next_mem/curr_mem explanation.
    data_table = doc.tables[2]
    replace_table_data(
        data_table,
        [
            ("字段类别", "字段", "作用"),
            ("会话上下文", "session_id, sender_id, epoch", "将消息绑定到具体会话、发送方和状态阶段，防止跨会话或跨阶段混淆。"),
            ("顺序状态", "seq", "保证消息单调递增，检测重放、旧包和序列缺口。"),
            ("内容绑定", "payload, payload_hash", "将应用层内容绑定到认证计算与记忆链更新中。"),
            ("链式状态", "prev_mem", "表示发送前的前序记忆状态，用于与接收端本地 last_mem 比较，验证历史连续性。"),
            ("链式结果", "next_mem / curr_mem（可选）", "表示发送方计算出的消息后记忆状态；接收端可重新计算并比对，适用于调试、审计或跨节点同步。"),
            ("时间辅助", "timestamp", "辅助检测过期、异常延迟或审计排序。"),
            ("认证字段", "auth_tag", "对会话上下文、顺序状态、内容绑定和链式状态等关键字段进行 HMAC 认证。"),
        ],
        [1.15, 2.25, 3.10],
    )
    table_note = insert_paragraph_after_table(
        data_table,
        "其中，prev_mem 是 DATA 报文中必需的链式状态输入，接收端必须验证 P.prev_mem == local.last_mem；随后接收端根据 payload_hash、"
        "seq、session_id、epoch 和 sender_id 重新计算 M_i。next_mem / curr_mem 属于可选扩展字段，在当前原型中可由接收端本地计算得到，"
        "因此不作为必需传输字段；若显式携带，接收端应重新计算 M_i' 并在 P.next_mem != M_i' 时拒绝该报文。",
    )
    format_paragraph(table_note)

    # DATA verification algorithm with optional next_mem check.
    alg1 = doc.tables[3]
    alg_lines = [
        "Input: packet P, expected_session_id, last_seq, last_mem, key K",
        "Output: accept/reject decision and updated memory state",
        '1: if required fields are missing then reject',
        '2: if P.type != DATA then reject',
        '3: if P.session_id != expected_session_id then reject',
        "4: h' = H(P.payload)",
        "5: if h' != P.payload_hash then reject",
        "6: tag' = HMAC(K, P without auth_tag)",
        "7: if tag' != P.auth_tag then reject",
        "8: if P.seq <= last_seq then reject as replay or stale packet",
        "9: if P.seq > last_seq + 1 then reject as sequence gap",
        "10: if P.prev_mem != last_mem then reject as history discontinuity",
        "11: M_i = H(last_mem || P.session_id || P.epoch || P.seq || P.payload_hash || P.sender_id)",
        "12: if P.next_mem is present and P.next_mem != M_i then reject as memory-result mismatch",
        "13: update last_seq = P.seq and last_mem = M_i",
        "14: accept P",
    ]
    while len(alg1.rows) > len(alg_lines):
        alg1._tbl.remove(alg1.rows[-1]._tr)
    while len(alg1.rows) < len(alg_lines):
        alg1.add_row()
    for row, text in zip(alg1.rows, alg_lines):
        set_cell_text(row.cells[0], text, bold=False, font_size=9.5)

    # Recovery flow: explicitly state what is recovered.
    recovery_step4 = find_paragraph(doc, lambda t: t.startswith("4. 验证通过后，服务器返回恢复状态"))
    set_text(
        recovery_step4,
        "4. 验证通过后，服务器返回恢复状态（last_seq、last_mem、checkpoint_seq、checkpoint_mem）；验证失败时返回明确拒绝原因，并拒绝更新状态。",
    )
    recovery_step5 = find_paragraph(doc, lambda t: t.startswith("5. 客户端从 last_seq"))
    set_text(
        recovery_step5,
        "5. 客户端从 last_seq + 1 继续通信，并用恢复后的 last_mem 作为后续 memory chain 的起点。",
    )
    recovery_note = insert_paragraph_after(
        recovery_step5,
        "因此，GMCP-R 的恢复对象不是应用层 payload 全文，也不是单纯网络连接句柄，而是能够作为后续安全通信起点的历史记忆状态："
        "last_seq 给出恢复点，last_mem 给出该恢复点对应的链式状态，Checkpoint 给出可验证的近端快照。三者一致时，通信双方才能继续推进新的 DATA 报文。",
    )
    format_paragraph(recovery_note)

    # Security analysis with Property / Proof Sketch style.
    msg_integrity = find_paragraph(doc, lambda t: t.startswith("每个 DATA 报文的 auth_tag"))
    set_text(
        msg_integrity,
        "Claim. Message Integrity. 在 HMAC-SHA256 满足不可伪造性的假设下，攻击者难以修改 DATA 报文中的 payload、seq、prev_mem 或会话上下文字段而仍使接收端接受该报文。",
    )
    p = insert_paragraph_after(
        msg_integrity,
        "Proof Sketch. 每个 DATA 报文的 auth_tag 覆盖除认证标签自身外的关键字段，包括 payload_hash、seq、prev_mem、session_id、epoch 与 sender_id。攻击者若修改任一字段，则服务器重新计算得到的 HMAC 输入发生变化。若修改后仍能通过验证，则意味着攻击者在未知共享密钥 K 的情况下伪造了有效 HMAC 标签，与 HMAC 不可伪造性假设矛盾。",
    )
    format_paragraph(p)

    history_para = find_paragraph(doc, lambda t: t.startswith("GMCP-R 的记忆状态按链式方式演进"))
    set_text(
        history_para,
        "Property 1. History Continuity. 在哈希函数 H 满足抗碰撞性，且 HMAC 不可伪造的假设下，攻击者难以构造一条与已接受历史不同的消息序列，却使接收端得到相同且可接受的 last_mem。",
    )
    p = insert_paragraph_after(
        history_para,
        "Proof Sketch. 设两条历史序列在某一位置 j 首次不同，但最终均被接收端接受并得到相同记忆状态 M_i。由于 M_i 递归依赖 M_{i-1}、session_id、epoch、seq_i、payload_hash_i 和 sender_id，位置 j 的差异会改变后续所有记忆状态输入。若最终状态仍相同，则要么存在一次哈希碰撞，要么攻击者伪造了覆盖差异字段的 HMAC 标签。根据 H 的抗碰撞性和 HMAC 的不可伪造性，该事件概率可忽略。因此，不同历史难以产生相同可接受记忆状态。",
    )
    format_paragraph(p)

    rollback_heading = find_paragraph(doc, lambda t: t == "5.3 重放与回滚抵抗")
    set_text(rollback_heading, "5.3 票据真实性与恢复后一致性")
    rollback_para = find_paragraph(doc, lambda t: t.startswith("报文级重放由单调递增的 seq 检测"))
    set_text(
        rollback_para,
        "Property 2. Ticket Unforgeability and Recovery Consistency. 在 HMAC 安全假设下，攻击者不能修改 MemoryTicket 中的 last_seq、last_mem、checkpoint、expire_time、nonce 或上下文字段而仍通过服务端验证；通过验证的票据必须描述同一会话、同一客户端、同一 epoch 下的一致恢复状态。",
    )
    p = insert_paragraph_after(
        rollback_para,
        "Proof Sketch. MemoryTicket 的 server_auth_tag 覆盖 sid、cid、epoch、last_seq、last_mem、ckpt_seq、ckpt_mem、exp、nonce 和 key_version。攻击者修改任一字段都会改变 HMAC 输入；若修改后仍通过验证，则等价于成功伪造 HMAC 标签。恢复算法还显式检查 session_id、client_id 与 epoch 是否匹配请求上下文，并检查 checkpoint_seq <= last_seq。因此，服务端只接受与当前恢复请求一致、且结构上合法的恢复状态。",
    )
    format_paragraph(p)

    ticket_heading = find_paragraph(doc, lambda t: t == "5.4 票据真实性与恢复后一致性")
    set_text(ticket_heading, "5.4 重放与回滚抵抗")
    ticket_para = find_paragraph(doc, lambda t: t.startswith("MemoryTicket 的 server_auth_tag 对票据字段整体签名"))
    set_text(
        ticket_para,
        "Property 3. Rollback Resistance. 若服务端维护最小可接受恢复序号 min_last_seq，并记录已消费 ticket_nonce，则任何 last_seq < min_last_seq 的旧票据或重复使用票据均会被拒绝。",
    )
    p = insert_paragraph_after(
        ticket_para,
        "Proof Sketch. 服务端在恢复算法中显式检查 T.last_seq < S.min_last_seq，并在该条件成立时返回 rollback detected；同时，已消费的 ticket_nonce 会被记录并用于拒绝重复恢复请求。因此，攻击者无法利用旧票据将服务端恢复到低于 min_last_seq 的历史位置，也无法通过重复提交同一有效票据绕过状态单调性检查。报文级重放还会被 seq <= last_seq 检测，避免恢复后接受旧 DATA 报文。",
    )
    format_paragraph(p)

    conclusion = find_paragraph(doc, lambda t: t.startswith("本文提出 GMCP-R，一种面向断续通信"))
    set_text(
        conclusion,
        "本文提出 GMCP-R，一种面向断续通信的历史连续性安全恢复协议。该协议通过 memory chain 绑定消息历史，通过 MemoryTicket 验证恢复请求，通过 Checkpoint 控制恢复开销；恢复过程恢复的是 last_seq、last_mem 与 checkpoint 共同描述的历史记忆状态，而不是单纯恢复一条网络连接，从而缓解“连接恢复不等于历史状态恢复”的问题。",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
