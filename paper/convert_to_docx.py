#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将Markdown论文转换为Word格式
"""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_document():
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    return doc

def add_title(doc, title, subtitle=None):
    # 添加标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        run.italic = True
        run.font.size = Pt(14)

def add_heading(doc, text, level=1):
    if level == 1:
        p = doc.add_heading(text, level=1)
    elif level == 2:
        p = doc.add_heading(text, level=2)
    elif level == 3:
        p = doc.add_heading(text, level=3)
    else:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True

def add_paragraph(doc, text):
    p = doc.add_paragraph(text)
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    # 添加表头
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    # 添加数据行
    for row in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            row_cells[i].text = str(cell_text)

def parse_markdown(md_file):
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    return content

def convert_md_to_docx(md_file, docx_file):
    doc = create_document()
    
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    in_table = False
    table_headers = []
    table_rows = []
    
    for line in lines:
        line = line.rstrip('\n')
        
        # 跳过空行
        if not line.strip():
            if in_table and table_headers:
                add_table(doc, table_headers, table_rows)
                table_headers = []
                table_rows = []
                in_table = False
            continue
        
        # 标题
        if line.startswith('# '):
            add_title(doc, line[2:])
        elif line.startswith('## '):
            add_heading(doc, line[3:], level=1)
        elif line.startswith('### '):
            add_heading(doc, line[4:], level=2)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], level=3)
        
        # 表格
        elif line.startswith('|'):
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            
            # 跳过分隔行
            if all(set(cell) <= {'-', ' ', ':'} for cell in cells):
                continue
            
            if not in_table:
                in_table = True
                table_headers = cells
            else:
                table_rows.append(cells)
        
        # 列表项
        elif line.startswith('- ') or line.startswith('* '):
            add_paragraph(doc, line)
        elif re.match(r'^\d+\. ', line):
            add_paragraph(doc, line)
        
        # 普通段落
        else:
            # 处理粗体和斜体
            p = doc.add_paragraph()
            
            # 简单处理：移除markdown格式标记
            clean_line = line.replace('**', '').replace('*', '').replace('`', '')
            p.add_run(clean_line)
    
    # 处理最后的表格
    if in_table and table_headers:
        add_table(doc, table_headers, table_rows)
    
    doc.save(docx_file)
    print(f"已生成Word文档: {docx_file}")

if __name__ == '__main__':
    md_file = '/Users/a0000/Desktop/实验/gmcp_r/paper/main_zh.md'
    docx_file = '/Users/a0000/Desktop/实验/gmcp_r/paper/GMCP-R论文.docx'
    
    convert_md_to_docx(md_file, docx_file)
