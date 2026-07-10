from __future__ import annotations

import csv
import math
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


INPUT = Path("/Users/a0000/Desktop/论文/GMCP-R_SCI中文初稿_内容修订版_2.7_按建议修订.docx")
OUTPUT = Path("/Users/a0000/Desktop/论文/GMCP-R_SCI中文初稿_内容修订版_2.8_补齐实验图表.docx")
ASSET_DIR = Path("/Users/a0000/Desktop/实验/gmcp_r/gmcpr_outline_assets")

REAL_RECOVERY = Path("results/real_recovery/summary_real_recovery.csv")
BASELINE_COST = Path("results/baseline/summary_baseline_comparison.csv")
MEMORY_CHAIN_FIG = Path("/Users/a0000/Desktop/实验/gmcp_r/gmcpr_revision_assets/gmcpr_memory_chain.png")


def load_font(size: int, mono: bool = False, bold: bool = False):
    candidates = []
    if mono:
        candidates += [
            "/System/Library/Fonts/Menlo.ttc",
            "/System/Library/Fonts/SFNSMono.ttf",
        ]
    if bold:
        candidates += [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
            "/System/Library/Fonts/PingFang.ttc",
        ]
    candidates += [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        try:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def draw_arrow(draw: ImageDraw.ImageDraw, start, end, color=(50, 75, 110), width=4):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - 16 * sign, y2 - 9), (x2 - 16 * sign, y2 + 9)]
    else:
        sign = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 9, y2 - 16 * sign), (x2 + 9, y2 - 16 * sign)]
    draw.polygon(pts, fill=color)


def centered_text(draw, box, text, font, fill=(20, 30, 44), line_gap=8):
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=fill)
        y += h + line_gap


def rounded_box(draw, box, text, font, fill, outline, text_fill=(20, 30, 44), radius=22):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=3)
    centered_text(draw, box, text, font, text_fill)


def create_data_packet_figure(path: Path):
    w, h = 1850, 680
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title = load_font(44, bold=True)
    font = load_font(25, mono=True)
    small = load_font(21, mono=True)
    draw.text((70, 45), "GMCP-R DATA Packet Format", font=title, fill=(21, 43, 72))
    draw.line((70, 105, w - 75, 105), fill=(185, 196, 210), width=2)

    fields = [
        ("type", 120, (229, 241, 255)),
        ("session_id", 190, (239, 246, 255)),
        ("sender_id", 170, (239, 246, 255)),
        ("epoch", 130, (239, 246, 255)),
        ("seq", 120, (244, 250, 246)),
        ("payload_hash", 230, (255, 248, 238)),
        ("prev_mem", 210, (255, 242, 242)),
        ("timestamp", 190, (250, 250, 250)),
        ("auth_tag", 210, (242, 244, 250)),
    ]
    x = 85
    y1, y2 = 205, 315
    for name, width, fill in fields:
        draw.rounded_rectangle((x, y1, x + width, y2), radius=12, fill=fill, outline=(75, 100, 135), width=3)
        centered_text(draw, (x, y1, x + width, y2), name, font)
        x += width + 10

    groups = [
        ("Session context", 215, 220, (239, 246, 255)),
        ("Sequence state", 745, 120, (244, 250, 246)),
        ("Payload binding", 875, 230, (255, 248, 238)),
        ("Chain state", 1115, 210, (255, 242, 242)),
        ("Authentication", 1535, 210, (242, 244, 250)),
    ]
    for label, gx, gw, color in groups:
        draw.line((gx, 178, gx + gw, 178), fill=(80, 105, 138), width=3)
        draw.line((gx, 178, gx, 193), fill=(80, 105, 138), width=3)
        draw.line((gx + gw, 178, gx + gw, 193), fill=(80, 105, 138), width=3)
        bbox = draw.textbbox((0, 0), label, font=small)
        draw.text((gx + (gw - (bbox[2] - bbox[0])) / 2, 142), label, font=small, fill=(45, 64, 90))

    rounded_box(
        draw,
        (250, 420, 1450, 565),
        "Receiver verification:\n1) HMAC over packet fields   2) seq == last_seq + 1   3) prev_mem == local.last_mem\n4) compute M_i locally; optional next_mem/curr_mem is accepted only if it matches M_i",
        small,
        fill=(248, 251, 255),
        outline=(112, 143, 178),
    )
    draw_arrow(draw, (700, 320), (700, 420), color=(80, 105, 138), width=4)
    img.save(path)


def create_recovery_flow_figure(path: Path):
    w, h = 1700, 920
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title = load_font(44, bold=True)
    head = load_font(32, bold=True)
    font = load_font(24, mono=True)
    small = load_font(22, mono=True)
    draw.text((70, 45), "MemoryTicket-Based Secure Recovery Procedure", font=title, fill=(21, 43, 72))
    draw.line((70, 105, 1625, 105), fill=(185, 196, 210), width=2)

    cx, sx = 285, 1285
    draw.text((cx - 75, 135), "Client", font=head, fill=(24, 55, 92))
    draw.text((sx - 75, 135), "Server", font=head, fill=(24, 55, 92))
    draw.line((cx, 185, cx, 820), fill=(205, 214, 226), width=4)
    draw.line((sx, 185, sx, 820), fill=(205, 214, 226), width=4)

    rounded_box(draw, (80, 220, 490, 330), "detect disconnect /\nattack / seq gap", font, (255, 248, 235), (197, 145, 65))
    draw_arrow(draw, (490, 385), (1060, 385), color=(176, 85, 85), width=5)
    rounded_box(draw, (80, 340, 490, 430), "RECOVERY_REQUEST\n+ MemoryTicket", font, (255, 242, 242), (176, 85, 85))
    rounded_box(
        draw,
        (1060, 260, 1620, 505),
        "verify server_auth_tag\ncheck expire_time\ncheck ticket_nonce\ncheck session/client/epoch\ncheck last_seq and checkpoint",
        small,
        (255, 248, 248),
        (176, 85, 85),
    )

    draw_arrow(draw, (1060, 610), (490, 610), color=(54, 112, 75), width=5)
    rounded_box(draw, (80, 560, 490, 670), "RECOVERY_RESPONSE\nlast_seq, last_mem,\nckpt_seq, ckpt_mem", small, (244, 250, 246), (75, 136, 91))
    rounded_box(
        draw,
        (1060, 560, 1620, 710),
        "accept only if recovered state\nbelongs to the same legal history chain;\notherwise reject with explicit reason",
        small,
        (244, 250, 246),
        (75, 136, 91),
    )
    rounded_box(draw, (560, 760, 1140, 855), "continue DATA from last_seq + 1\nusing recovered last_mem as chain start", small, (239, 246, 255), (72, 117, 180))
    draw_arrow(draw, (490, 705), (560, 800), color=(72, 117, 180), width=4)
    draw_arrow(draw, (1060, 705), (1140, 800), color=(72, 117, 180), width=4)
    img.save(path)


def read_csv_dict(path: Path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def row_float(row, *names, default=0.0):
    for name in names:
        value = row.get(name)
        if value not in (None, ""):
            return float(value)
    return default


def row_percent(row, *names, default=0.0):
    value = row_float(row, *names, default=default)
    return value if value > 1 else value * 100


def bar_chart(draw, x, y, width, height, labels, series, colors, y_max, y_label, font, small, percent=False):
    draw.rectangle((x, y, x + width, y + height), outline=(80, 90, 105), width=2)
    # grid
    for i in range(1, 5):
        gy = y + height - height * i / 5
        draw.line((x, gy, x + width, gy), fill=(225, 230, 236), width=1)
    group_w = width / len(labels)
    n = len(series)
    bar_w = min(28, group_w / (n + 1.5))
    for li, label in enumerate(labels):
        gx = x + li * group_w + group_w * 0.15
        for si, (_, values) in enumerate(series):
            v = values[li]
            bh = height * (v / y_max if y_max else 0)
            bx = gx + si * (bar_w + 8)
            draw.rectangle((bx, y + height - bh, bx + bar_w, y + height), fill=colors[si])
        bbox = draw.textbbox((0, 0), label, font=small)
        draw.text((x + li * group_w + (group_w - (bbox[2] - bbox[0])) / 2, y + height + 12), label, font=small, fill=(25, 35, 48))
    # legend
    lx, ly = x + 10, y - 36
    for si, (name, _) in enumerate(series):
        draw.rectangle((lx, ly + si * 25, lx + 18, ly + 18 + si * 25), fill=colors[si])
        draw.text((lx + 26, ly - 2 + si * 25), name, font=small, fill=(25, 35, 48))
    draw.text((x - 10, y - 72), y_label, font=font, fill=(25, 35, 48))


def create_recovery_results_figure(path: Path):
    rows = read_csv_dict(REAL_RECOVERY)
    order = ["drop", "modify", "replay", "prev_mem", "disconnect"]
    by = {r["attack_type"]: r for r in rows}
    labels = order
    succ = [row_percent(by[k], "recovery_success_rate") for k in labels]
    mem = [row_percent(by[k], "memory_match_rate", "memory_match_after_recovery") for k in labels]
    seq = [row_percent(by[k], "final_seq_consistency_rate", "final_seq_consistency") for k in labels]
    latency = [row_float(by[k], "latency_mean_ms", "recovery_latency_ms") for k in labels]

    img = Image.new("RGB", (1700, 900), "white")
    draw = ImageDraw.Draw(img)
    title = load_font(42, bold=True)
    font = load_font(26, bold=True)
    small = load_font(21)
    draw.text((70, 45), "Secure Recovery after Attacks and Disconnections", font=title, fill=(21, 43, 72))
    draw.line((70, 105, 1625, 105), fill=(185, 196, 210), width=2)
    bar_chart(
        draw,
        120,
        190,
        690,
        470,
        labels,
        [("recovery", succ), ("memory", mem), ("seq", seq)],
        [(70, 120, 180), (80, 155, 96), (230, 148, 55)],
        110,
        "Rates (%)",
        font,
        small,
        percent=True,
    )
    # Latency panel
    x, y, width, height = 940, 190, 620, 470
    draw.rectangle((x, y, x + width, y + height), outline=(80, 90, 105), width=2)
    max_lat = max(latency) * 1.15
    for i in range(1, 5):
        gy = y + height - height * i / 5
        draw.line((x, gy, x + width, gy), fill=(225, 230, 236), width=1)
    group_w = width / len(labels)
    for i, label in enumerate(labels):
        v = latency[i]
        bh = height * v / max_lat
        bx = x + i * group_w + group_w * 0.28
        bw = group_w * 0.44
        draw.rectangle((bx, y + height - bh, bx + bw, y + height), fill=(93, 132, 180))
        draw.text((bx + 4, y + height - bh - 28), f"{v:.0f}", font=small, fill=(25, 35, 48))
        bbox = draw.textbbox((0, 0), label, font=small)
        draw.text((x + i * group_w + (group_w - (bbox[2] - bbox[0])) / 2, y + height + 12), label, font=small, fill=(25, 35, 48))
    draw.text((x - 10, y - 45), "Recovery latency (ms)", font=font, fill=(25, 35, 48))
    rounded_box(
        draw,
        (145, 730, 1560, 835),
        "All tested scenarios recovered successfully with memory match and final sequence consistency = 100%.\nExtra recovery messages are fixed at 2; extra bytes stay around 1.37-1.39 KB.",
        small,
        (248, 251, 255),
        (112, 143, 178),
    )
    img.save(path)


def create_checkpoint_cost_figure(path: Path):
    rows = read_csv_dict(BASELINE_COST)
    order = ["gmcp_r", "hash_chain", "seq_mac", "ticket_only"]
    by = {r["protocol"]: r for r in rows}
    labels = ["GMCP-R", "Hash Chain", "Seq+MAC", "Ticket Only"]
    replay = [float(by[k]["replay_count_mean"]) for k in order]
    bytes_ = [float(by[k]["recovery_extra_bytes_mean"]) for k in order]

    img = Image.new("RGB", (1700, 840), "white")
    draw = ImageDraw.Draw(img)
    title = load_font(42, bold=True)
    font = load_font(26, bold=True)
    small = load_font(21)
    draw.text((70, 45), "Recovery Cost and Checkpoint Effectiveness", font=title, fill=(21, 43, 72))
    draw.line((70, 105, 1625, 105), fill=(185, 196, 210), width=2)

    def horizontal_panel(x, y, w, h, values, label, color, log=False, suffix=""):
        draw.text((x, y - 44), label, font=font, fill=(25, 35, 48))
        draw.rectangle((x, y, x + w, y + h), outline=(80, 90, 105), width=2)
        scaled = [math.log10(v + 1) if log else v for v in values]
        mx = max(scaled) * 1.08
        row_h = h / len(values)
        for i, name in enumerate(labels):
            yy = y + i * row_h + row_h * 0.25
            bw = (w - 210) * scaled[i] / mx if mx else 0
            draw.text((x + 14, yy + 6), name, font=small, fill=(25, 35, 48))
            draw.rectangle((x + 155, yy, x + 155 + bw, yy + row_h * 0.48), fill=color)
            val = values[i]
            txt = f"{val:,.0f}{suffix}" if val >= 10 else f"{val:.1f}{suffix}"
            draw.text((x + 165 + bw, yy + 4), txt, font=small, fill=(25, 35, 48))

    horizontal_panel(125, 205, 650, 420, replay, "Replay count needed for recovery", (70, 120, 180), log=False)
    horizontal_panel(925, 205, 650, 420, bytes_, "Extra recovery bytes (log scale)", (220, 126, 68), log=True, suffix=" B")
    rounded_box(
        draw,
        (155, 700, 1545, 790),
        "Checkpoint limits GMCP-R recovery to a bounded suffix, while pure Hash Chain must replay far more history.",
        small,
        (248, 251, 255),
        (112, 143, 178),
    )
    img.save(path)


def apply_run_font(run, size: float | None = None, bold: bool | None = None):
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "Songti SC")
    r_fonts.set(qn("w:cs"), "Times New Roman")


def set_text(paragraph, text: str, size: float | None = None, bold: bool | None = None):
    paragraph.clear()
    run = paragraph.add_run(text)
    apply_run_font(run, size=size, bold=bold)
    return paragraph


def add_text(paragraph, text: str, size: float | None = None, bold: bool | None = None):
    run = paragraph.add_run(text)
    apply_run_font(run, size=size, bold=bold)
    return run


def find_para(doc: Document, predicate):
    for p in doc.paragraphs:
        if predicate(p.text.strip()):
            return p
    raise ValueError("paragraph not found")


def insert_para_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        add_text(p, text)
    return p


def insert_para_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        add_text(p, text)
    return p


def insert_after_table(table, text="", style=None):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    from docx.text.paragraph import Paragraph

    p = Paragraph(new_p, table._parent)
    if style:
        p.style = style
    if text:
        add_text(p, text)
    return p


def add_caption_after(anchor_para, caption: str):
    p = insert_para_after(anchor_para, caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        apply_run_font(r, size=10)
    return p


def add_picture_after(anchor_para, image_path: Path, caption: str, width=6.3):
    p = insert_para_after(anchor_para)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    return add_caption_after(p, caption)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def resize_drawings_in_paragraph(paragraph, width_in: float):
    width_emu = int(Inches(width_in))
    for inline in paragraph._p.xpath('.//*[local-name()="inline"]'):
        extent = inline.find(qn("wp:extent"))
        if extent is None:
            continue
        old_cx = int(extent.get("cx"))
        old_cy = int(extent.get("cy"))
        height_emu = int(old_cy * width_emu / old_cx)
        extent.set("cx", str(width_emu))
        extent.set("cy", str(height_emu))
        for graphic_extent in inline.xpath('.//*[local-name()="ext"]'):
            if graphic_extent.get("cx") == str(old_cx) and graphic_extent.get("cy") == str(old_cy):
                graphic_extent.set("cx", str(width_emu))
                graphic_extent.set("cy", str(height_emu))


def replace_figure_2(doc: Document, image_path: Path):
    paragraphs = list(doc.paragraphs)
    for caption_index, caption in enumerate(paragraphs):
        if caption.text.strip() != "图 2. GMCP-R 的记忆链状态演化过程":
            continue
        anchor = find_para(doc, lambda t: t.startswith("图 2 展示了 GMCP-R 中 memory chain"))
        image_index = None
        for idx in range(caption_index - 1, -1, -1):
            if paragraphs[idx]._p.xpath('.//*[local-name()="drawing"]'):
                image_index = idx
                break
        if image_index is None:
            return
        for old_para in reversed(paragraphs[image_index : caption_index + 1]):
            delete_paragraph(old_para)
        image_para = insert_para_after(anchor)
        image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_para.paragraph_format.page_break_before = True
        image_para.paragraph_format.keep_with_next = True
        image_para.paragraph_format.space_before = Pt(0)
        image_para.paragraph_format.space_after = Pt(2)
        image_para.add_run().add_picture(str(image_path), width=Inches(5.8))
        caption = add_caption_after(image_para, "图 2. GMCP-R 的记忆链状态演化过程")
        caption.paragraph_format.keep_with_next = True
        return


def add_table_after(doc: Document, anchor, rows, widths=None, style="Table Grid"):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = style
    table.autofit = False
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 or len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, text, size=9.2, bold=(ri == 0))
            if widths:
                set_cell_width(cell, widths[ci])
    set_repeat_header(table)
    if hasattr(anchor, "_p"):
        anchor._p.addnext(table._tbl)
    else:
        anchor._tbl.addnext(table._tbl)
    return table


def set_cell_width(cell, width_in):
    width = Inches(width_in)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width.twips)))


def set_repeat_header(table):
    if not table.rows:
        return
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def remove_trailing_empty_table_rows(doc: Document):
    for table in doc.tables:
        while table.rows and all(not cell.text.strip() for cell in table.rows[-1].cells):
            table._tbl.remove(table.rows[-1]._tr)


def format_caption_paragraphs(doc: Document):
    for p in doc.paragraphs:
        t = p.text.strip()
        is_caption = (
            re.match(r"^图\s+\d+[.．]", t)
            or re.match(r"^表\s+\d+[.．]", t)
            or re.match(r"^附录图\s+[A-Z]\d+[.．]", t)
            or re.match(r"^附录表\s+[A-Z]\d+[.．]", t)
        )
        if is_caption:
            try:
                p.style = doc.styles["Normal"]
            except Exception:
                pass
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                apply_run_font(r, size=10, bold=t.startswith(("表 ", "附录表")))
        elif t.startswith(("图 ", "表 ")):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                apply_run_font(r, size=10.5, bold=False)


def fix_layout_issues(doc: Document):
    for p in doc.paragraphs:
        if p.text.strip() in {
            "表 2. GMCP-R 的主要状态变量",
            "表 9. 攻击后安全恢复结果",
            "表 10. MemoryTicket 无效票据拒绝结果",
        }:
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True


def replace_visible_numbering(doc: Document):
    replacements = {
        "表 4. MemoryTicket 核心字段与安全作用": "表 6. MemoryTicket 核心字段与安全作用",
        "表 5. 实验环境与参数概览": "表 7. 实验环境与参数概览",
        "表 6. Baseline 协议对比结果": "表 8. Baseline 协议对比结果",
        "表 7. MemoryTicket 无效票据拒绝结果": "表 10. MemoryTicket 无效票据拒绝结果",
        "表 8. 性能基准测试结果（本地回环，64B-1024B 均值）": "表 12. 性能基准测试结果（本地回环，64B-1024B 均值）",
        "表 9. 并发客户端测试结果": "表 13. 并发客户端测试结果",
        "图 3. 各协议攻击检测率与误接受率对比。": "图 5. 各协议攻击检测率与误接受率对比。",
        "图 4. 正常通信条件下各协议平均吞吐量对比。": "图 6. 正常通信条件下各协议平均吞吐量对比。",
        "图 5. 正常通信条件下各协议平均 RTT 对比。": "图 7. 正常通信条件下各协议平均 RTT 对比。",
        "图 6. 本地纯计算吞吐量随载荷大小变化。": "图 10. 本地纯计算吞吐量随载荷大小变化。",
        "图 7. 并发客户端数量与总吞吐量变化。": "图 11. 并发客户端数量与总吞吐量变化。",
        "图 8. 并发客户端数量与平均 RTT 变化。": "图 12. 并发客户端数量与平均 RTT 变化。",
    }
    inline_replacements = {
        "表 6 汇总": "表 8 汇总",
        "与表 6 一致": "与表 8 一致",
        "表 6 里": "表 8 里",
        "7.2 Invalid MemoryTicket Rejection": "7.3 Invalid MemoryTicket Rejection",
        "7.3 Performance Benchmark": "7.5 Performance Benchmark",
        "7.4 Concurrent Client Evaluation": "7.6 Concurrent Client Evaluation",
        "7.5 Preliminary Weak-Network Simulation": "7.7 Preliminary Weak-Network Simulation",
    }
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in replacements:
            set_text(p, replacements[t])
        else:
            new_t = p.text
            for old, new in inline_replacements.items():
                new_t = new_t.replace(old, new)
            if new_t != p.text:
                set_text(p, new_t)


def update_existing_sections(doc: Document):
    # 4.4
    p44 = find_para(doc, lambda t: t.startswith("Checkpoint 是周期性状态快照"))
    set_text(
        p44,
        "Checkpoint 是周期性状态快照，记录服务器已经验证并接受的近端历史状态。默认配置下，每处理 100 条 DATA 报文创建一个 Checkpoint；"
        "快照字段包括 session_id、epoch、ckpt_seq、ckpt_mem、created_at 和 checkpoint_auth_tag。恢复时，服务器可以从最近可信 Checkpoint "
        "开始验证或重放后续少量消息，而不必从会话初始状态重建全部历史。因此，Checkpoint 将纯 Hash Chain 的恢复成本从 O(n) 限制为 O(k)，"
        "其中 n 为会话历史长度，k 为 checkpoint 间隔或最近快照之后的待验证后缀长度。",
    )
    cap = insert_para_after(p44, "表 5. Checkpoint 核心字段与作用")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        apply_run_font(r, size=10, bold=True)
    tbl = add_table_after(
        doc,
        cap,
        [
            ["字段", "作用"],
            ["session_id / epoch", "将快照绑定到具体会话和状态阶段，避免跨会话或跨 epoch 恢复。"],
            ["ckpt_seq", "记录快照覆盖到的最新消息序号，作为恢复锚点的序列位置。"],
            ["ckpt_mem", "记录 ckpt_seq 对应的记忆链状态，作为后续恢复验证起点。"],
            ["created_at", "辅助判断快照新旧和审计恢复过程。"],
            ["checkpoint_auth_tag", "服务器对快照字段计算认证标签，防止快照被篡改。"],
        ],
        widths=[2.3, 4.1],
    )
    note = insert_after_table(
        tbl,
        "在恢复算法中，Checkpoint 不单独决定恢复成功，而是与 MemoryTicket 中的 last_seq 和 last_mem 联合使用。若 ckpt_seq > last_seq，"
        "或 ckpt_mem 无法作为 last_mem 所属历史链的可信前缀，服务端应拒绝恢复请求。",
    )
    for r in note.runs:
        apply_run_font(r, size=10.5)

    # 4.5
    p45 = find_para(doc, lambda t: t.startswith("MemoryTicket 的认证标签覆盖恢复相关字段"))
    set_text(
        p45,
        "MemoryTicket 由服务器在正常通信或安全恢复后签发，用于把恢复点、恢复状态和恢复锚点绑定为一个短期可验证凭证。"
        "票据不是简单的会话恢复 token，而是包含 last_seq、last_mem、ckpt_seq、ckpt_mem、expire_time、ticket_nonce 与 key_version 的状态承诺。",
    )
    formula = find_para(doc, lambda t: t.startswith("tag_T = HMAC"))
    set_text(formula, "tag_T = HMAC_{K_s}(Enc(sid, cid, epoch, last_seq, last_mem, ckpt_seq, ckpt_mem, exp, nonce, key_version))")
    after_formula = insert_para_after(
        formula,
        "其中，K_s 为服务器侧票据认证密钥。expire_time 限制票据有效期，ticket_nonce 用于一次性消费检测，key_version 支持服务器侧密钥轮换。"
        "服务端验证恢复请求时必须重新计算 tag_T，并同时检查 nonce 是否已使用、会话上下文是否匹配、last_seq 是否满足最小可接受恢复序号，以及 checkpoint 是否与恢复状态结构一致。",
    )
    for r in after_formula.runs:
        apply_run_font(r, size=10.5)

    # 5.3 / 5.5
    h53 = find_para(doc, lambda t: t.startswith("5.3 "))
    set_text(h53, "5.3 Ticket Unforgeability（票据不可伪造）")
    prop2 = find_para(doc, lambda t: t.startswith("Property 2. Ticket Unforgeability"))
    set_text(
        prop2,
        "Property 2. Ticket Unforgeability. 在 HMAC 安全假设下，攻击者不能修改 MemoryTicket 中的 last_seq、last_mem、checkpoint、expire_time、nonce 或上下文字段而仍通过服务端验证。",
    )
    proof2 = find_para(doc, lambda t: t.startswith("Proof Sketch. MemoryTicket 的 server_auth_tag"))
    set_text(
        proof2,
        "Proof Sketch. MemoryTicket 的 server_auth_tag 覆盖 sid、cid、epoch、last_seq、last_mem、ckpt_seq、ckpt_mem、exp、nonce 和 key_version。攻击者修改任一字段都会改变 HMAC 输入；若修改后仍通过验证，则等价于成功伪造 HMAC 标签，与 HMAC 不可伪造性假设矛盾。",
    )
    marker = find_para(doc, lambda t: t.startswith("需要强调的是"))
    h55 = insert_para_before(marker, "5.5 Recovery State Consistency（恢复状态一致性）", style="Heading 2")
    p55 = insert_para_after(
        h55,
        "Property 4. Recovery State Consistency. 若服务端恢复算法接受一个恢复请求，则返回的 last_seq、last_mem、checkpoint_seq 和 checkpoint_mem 必须属于同一会话、同一 epoch 下的一条合法历史链；服务端不会接受不属于合法历史链的恢复状态。",
    )
    p55b = insert_para_after(
        p55,
        "Proof Sketch. 服务端首先验证 MemoryTicket 的 HMAC，因此攻击者不能单独替换 last_seq、last_mem 或 checkpoint 字段。随后，恢复算法检查 session_id、client_id 与 epoch 是否匹配请求上下文，检查 checkpoint_seq <= last_seq，并检查 last_seq 不低于服务端维护的 min_last_seq。由于 last_mem 是由 memory chain 递归计算得到的状态，且 checkpoint_mem 是该链上的可信近端快照，任何与合法历史链不一致的恢复状态都会在票据认证、上下文绑定、checkpoint 结构检查或状态单调性检查中失败。因此，被接受的恢复状态可以作为后续 DATA 报文的安全起点。",
    )
    for p in (p55, p55b):
        for r in p.runs:
            apply_run_font(r, size=10.5)


def insert_design_figures(doc: Document, data_packet_fig: Path, recovery_fig: Path):
    data_caption_anchor = find_para(doc, lambda t: t.startswith("表 4. GMCP-R DATA 报文字段"))
    intro = insert_para_before(
        data_caption_anchor,
        "图 3 给出了 DATA 报文中字段分组与验证关系。该图强调 prev_mem 属于链式状态输入，而 auth_tag 需要覆盖会话上下文、顺序状态、内容绑定和链式状态等关键字段。",
    )
    for r in intro.runs:
        apply_run_font(r, size=10.5)
    add_picture_after(intro, data_packet_fig, "图 3. GMCP-R DATA 报文格式与验证字段。", width=5.9)

    step5 = find_para(doc, lambda t: t.startswith("5. 客户端从 last_seq"))
    flow_intro = insert_para_after(
        step5,
        "图 4 展示了基于 MemoryTicket 的安全恢复流程。该流程的关键在于服务器先验证票据真实性和状态单调性，再返回恢复后的 last_seq 与 last_mem；若任一检查失败，恢复请求被显式拒绝。",
    )
    for r in flow_intro.runs:
        apply_run_font(r, size=10.5)
    add_picture_after(flow_intro, recovery_fig, "图 4. 基于 MemoryTicket 的安全恢复流程。", width=6.4)


def insert_recovery_semantics(doc: Document):
    alg2_table = None
    for table in doc.tables:
        if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip().startswith("Input: recovery request R"):
            alg2_table = table
            break
    if alg2_table is None:
        raise ValueError("Algorithm 2 table not found")
    h = insert_after_table(alg2_table, "4.7 Secure Recovery Semantics（安全恢复语义）", style="Heading 2")
    p1 = insert_para_after(
        h,
        "GMCP-R 中“恢复”的语义不是恢复 TCP 连接句柄，也不是恢复完整 payload 历史，而是恢复可作为后续安全通信起点的历史记忆状态。"
        "last_seq 决定通信双方从哪个消息序号之后继续，last_mem 决定该恢复点对应的链式历史状态，Checkpoint 则限制恢复验证或历史重放的范围。",
    )
    p2 = insert_para_after(
        p1,
        "一次恢复被接受，需要满足三个条件：第一，MemoryTicket 能通过服务器认证标签验证，说明恢复状态字段未被篡改；第二，"
        "last_seq、last_mem、ckpt_seq 和 ckpt_mem 在结构上相互一致，说明恢复状态可解释为同一条合法历史链上的状态；第三，"
        "nonce、expire_time 和 min_last_seq 等检查通过，说明该恢复请求不是旧票据重放或历史回滚。只有三者同时成立，客户端才允许从 last_seq + 1 继续发送 DATA 报文。",
    )
    for p in (p1, p2):
        for r in p.runs:
            apply_run_font(r, size=10.5)


def insert_experiment_setup_sections(doc: Document):
    p62 = find_para(doc, lambda t: t.startswith("真实 baseline 对比实验覆盖"))
    set_text(
        p62,
        "本文设置三类 baseline：Hash Chain、Seq+MAC 和 Ticket Only。Hash Chain 用于比较纯链式历史验证能力，Seq+MAC 用于比较传统消息认证与序列号保护，Ticket Only 用于比较仅依赖恢复凭证的方案。真实 baseline 对比实验覆盖 4 种协议、5 类攻击类型、3 种消息数量、2 种载荷大小与 3 次重复，共 360 条记录。",
    )
    h63 = insert_para_after(p62, "6.3 Attack and Recovery Scenarios（攻击与恢复场景）", style="Heading 2")
    p63 = insert_para_after(
        h63,
        "攻击场景包括 drop、modify、replay 和 prev_mem。drop 用于模拟消息缺失与序列号间隙；modify 用于模拟 payload 或认证字段篡改；"
        "replay 用于模拟旧消息重新注入；prev_mem 用于直接检验历史连续性字段是否被篡改。恢复场景覆盖 drop recovery、modify recovery、"
        "replay recovery、prev_mem recovery 和 disconnect recovery。票据攻击另行覆盖 expired_ticket、replayed_ticket、tampered_ticket、rollback_ticket 和 wrong_session_ticket 五类无效 MemoryTicket。",
    )
    h64 = insert_para_after(p63, "6.4 Metrics（评价指标）", style="Heading 2")
    p64 = insert_para_after(
        h64,
        "主要指标包括攻击检测率、误接受率、误拒绝率、恢复成功率、恢复后 memory 一致率、最终序列一致性、恢复延迟、额外恢复消息数、额外恢复字节数、吞吐量、RTT 和并发成功率。对于代码级弱网仿真，本文只将其作为恢复行为的补充观察，不将其作为真实网络损伤下的核心结论。",
    )
    for p in (p63, p64):
        for r in p.runs:
            apply_run_font(r, size=10.5)


def recovery_table_rows():
    rows = [["场景", "攻击检测率", "恢复成功率", "memory 一致率", "seq 一致率", "恢复延迟(ms)", "额外消息", "额外字节"]]
    order = ["drop", "modify", "replay", "prev_mem", "disconnect"]
    names = {"drop": "drop", "modify": "modify", "replay": "replay", "prev_mem": "prev_mem", "disconnect": "disconnect"}
    by = {r["attack_type"]: r for r in read_csv_dict(REAL_RECOVERY)}
    for k in order:
        r = by[k]
        rows.append(
            [
                names[k],
                f'{row_percent(r, "attack_detection_rate"):.0f}%',
                f'{row_percent(r, "recovery_success_rate"):.0f}%',
                f'{row_percent(r, "memory_match_rate", "memory_match_after_recovery"):.0f}%',
                f'{row_percent(r, "final_seq_consistency_rate", "final_seq_consistency"):.0f}%',
                f'{row_float(r, "latency_mean_ms", "recovery_latency_ms"):.1f}',
                f'{row_float(r, "recovery_extra_messages_mean", "recovery_extra_messages"):.0f}',
                f'{row_float(r, "recovery_extra_bytes_mean", "recovery_extra_bytes"):.0f}',
            ]
        )
    return rows


def checkpoint_cost_rows():
    rows = [["协议", "恢复延迟(ms)", "memory 恢复率", "安全快速恢复率", "额外字节", "重放消息数"]]
    order = ["gmcp_r", "hash_chain", "seq_mac", "ticket_only"]
    labels = {"gmcp_r": "GMCP-R", "hash_chain": "Hash Chain", "seq_mac": "Seq+MAC", "ticket_only": "Ticket Only"}
    by = {r["protocol"]: r for r in read_csv_dict(BASELINE_COST)}
    for k in order:
        r = by[k]
        rows.append(
            [
                labels[k],
                f'{float(r["recovery_latency_ms_mean"]):.2f}',
                f'{float(r["memory_recovery_rate"]) * 100:.0f}%',
                f'{float(r["fast_secure_memory_recovery_rate"]) * 100:.0f}%',
                f'{float(r["recovery_extra_bytes_mean"]):,.0f}',
                f'{float(r["replay_count_mean"]):.1f}',
            ]
        )
    return rows


def insert_results_sections(doc: Document, recovery_results_fig: Path, checkpoint_fig: Path):
    # Renumber old headings before inserting new sections. These may already
    # have been updated by replace_visible_numbering(), so tolerate both forms.
    for old_prefix, new_text in [
        ("7.2 Invalid MemoryTicket", "7.3 Invalid MemoryTicket Rejection"),
        ("7.3 Invalid MemoryTicket", "7.3 Invalid MemoryTicket Rejection"),
        ("7.3 Performance Benchmark", "7.5 Performance Benchmark"),
        ("7.5 Performance Benchmark", "7.5 Performance Benchmark"),
        ("7.4 Concurrent Client Evaluation", "7.6 Concurrent Client Evaluation"),
        ("7.6 Concurrent Client Evaluation", "7.6 Concurrent Client Evaluation"),
        ("7.5 Preliminary Weak-Network Simulation", "7.7 Preliminary Weak-Network Simulation"),
        ("7.7 Preliminary Weak-Network Simulation", "7.7 Preliminary Weak-Network Simulation"),
    ]:
        try:
            set_text(find_para(doc, lambda t, old_prefix=old_prefix: t.startswith(old_prefix)), new_text)
        except ValueError:
            pass

    invalid_h = find_para(doc, lambda t: t.startswith("7.3 Invalid MemoryTicket"))
    h72 = insert_para_before(invalid_h, "7.2 Secure Recovery after Attacks and Disconnections", style="Heading 2")
    p72 = insert_para_after(
        h72,
        "本节直接回答“是不是恢复了”的问题。实验在真实 TCP 原型上构造 drop、modify、replay、prev_mem 和 disconnect 五类异常场景。"
        "客户端在检测到攻击、序列缺口或断连后发起 RECOVERY_REQUEST，服务器基于 MemoryTicket 和 Checkpoint 返回恢复状态，并要求客户端从 last_seq + 1 继续通信。",
    )
    cap = insert_para_after(p72, "表 9. 攻击后安全恢复结果")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        apply_run_font(r, size=10, bold=True)
    tbl = add_table_after(doc, cap, recovery_table_rows(), widths=[0.9, 0.85, 0.85, 0.9, 0.85, 0.9, 0.75, 0.85])
    src = insert_after_table(tbl, "数据来源：results/real_recovery/summary_real_recovery.csv。")
    for r in src.runs:
        apply_run_font(r, size=9)
    p_after = insert_para_after(
        src,
        "结果显示，五类场景下 GMCP-R 的攻击检测率、恢复成功率、恢复后 memory 一致率和最终序列一致率均为 100%。恢复过程平均只额外使用 2 条消息，额外字节约 1.37–1.39 KB。disconnect 场景恢复延迟较高，主要反映连接中断后的等待和重连开销；drop、modify、replay 与 prev_mem 场景的恢复延迟集中在约 174–180 ms。",
    )
    for r in p_after.runs:
        apply_run_font(r, size=10.5)
    add_picture_after(p_after, recovery_results_fig, "图 8. 攻击和断连后的安全恢复结果。", width=6.4)

    perf_h = find_para(doc, lambda t: t.startswith("7.5 Performance Benchmark"))
    h74 = insert_para_before(perf_h, "7.4 Recovery Cost and Checkpoint Effectiveness", style="Heading 2")
    p74 = insert_para_after(
        h74,
        "本节用于说明 Checkpoint 的价值。由于真实网络恢复延迟受公网抖动影响较大，本文同时报告协议级恢复成本模型，用于比较不同协议在恢复时需要重放的历史长度和额外字节数。该结果用于解释机制开销，不替代真实 TCP 端到端性能结论。",
    )
    cap2 = insert_para_after(p74, "表 11. 恢复成本与 Checkpoint 效果对比")
    cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap2.runs:
        apply_run_font(r, size=10, bold=True)
    tbl2 = add_table_after(doc, cap2, checkpoint_cost_rows(), widths=[1.05, 1.0, 1.0, 1.15, 1.05, 1.05])
    src2 = insert_after_table(tbl2, "数据来源：results/baseline/summary_baseline_comparison.csv（protocol-level simulation）。")
    for r in src2.runs:
        apply_run_font(r, size=9)
    p74b = insert_para_after(
        src2,
        "可以看到，GMCP-R 在保持 100% memory 恢复率和 100% 安全快速恢复率的同时，将平均重放消息数限制在约 63 条，额外恢复字节约 6.6 KB；纯 Hash Chain 需要重放约 1,155 条消息，额外字节约 369.6 KB。该差异体现了 Checkpoint 将恢复成本限制在近端后缀而非完整历史上的作用。",
    )
    for r in p74b.runs:
        apply_run_font(r, size=10.5)
    add_picture_after(p74b, checkpoint_fig, "图 9. Checkpoint 对恢复重放成本和额外字节的影响。", width=6.4)


def update_discussion(doc: Document):
    h81 = find_para(doc, lambda t: t.startswith("8.1 "))
    set_text(h81, "8.1 GMCP-R 实际解决的问题")
    ps = doc.paragraphs
    idx = next(i for i, p in enumerate(ps) if p._p is h81._p)
    set_text(
        ps[idx + 1],
        "GMCP-R 的目标不是替代 TLS/QUIC，也不是成为吞吐量最高的轻量认证协议，而是解决断续通信恢复后的历史状态可信性问题。"
        "实验中，GMCP-R 与 Hash Chain 均能检测历史断裂，但 GMCP-R 进一步通过 MemoryTicket 和 Checkpoint 给出可验证、可恢复且成本有界的恢复层。",
    )
    h82 = insert_para_after(ps[idx + 1], "8.2 恢复不只是重连", style="Heading 2")
    p82 = insert_para_after(
        h82,
        "普通重连只说明双方重新获得了传输通道；GMCP-R 的恢复还要求 last_seq、last_mem 与 checkpoint 能够对应同一条合法历史链。"
        "因此，恢复成功意味着客户端可以从 last_seq + 1 接续发送新 DATA 报文，且接收端能够用恢复后的 last_mem 继续验证后续历史连续性。",
    )
    for r in p82.runs:
        apply_run_font(r, size=10.5)
    set_text(find_para(doc, lambda t: t.startswith("8.2 与纯 Hash Chain")), "8.3 安全性与性能的权衡")
    para = find_para(doc, lambda t: t.startswith("实验中 GMCP-R 与 Hash Chain"))
    set_text(
        para,
        "GMCP-R 的吞吐量低于 Ticket Only、Seq+MAC 和纯 Hash Chain，主要原因是每条消息需要同时计算 payload_hash、更新 memory state、覆盖更多上下文字段的 HMAC，并维护恢复票据和 checkpoint 状态。该额外开销换来的能力是恢复后历史状态是否可信，以及恢复成本是否能被 Checkpoint 限制。对于工业遥测、审计日志、金融流水和医疗记录等高价值场景，这一安全目标通常比单纯追求最高吞吐量更重要。",
    )
    set_text(find_para(doc, lambda t: t.startswith("8.3 适用场景")), "8.4 适用场景")
    set_text(find_para(doc, lambda t: t.startswith("8.4 Resource Constraints")), "8.5 Resource Constraints in IoT and Edge Devices")
    set_text(find_para(doc, lambda t: t.startswith("8.5 Limitations")), "8.6 Limitations（局限性）")
    set_text(find_para(doc, lambda t: t.startswith("8.6 Future Work")), "8.7 Future Work: Post-Quantum Key Establishment Integration")


def main():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    data_packet_fig = ASSET_DIR / "figure_data_packet_format.png"
    recovery_flow_fig = ASSET_DIR / "figure_secure_recovery_flow.png"
    recovery_results_fig = ASSET_DIR / "figure_real_recovery_results.png"
    checkpoint_fig = ASSET_DIR / "figure_checkpoint_cost.png"
    create_data_packet_figure(data_packet_fig)
    create_recovery_flow_figure(recovery_flow_fig)
    create_recovery_results_figure(recovery_results_fig)
    create_checkpoint_cost_figure(checkpoint_fig)

    doc = Document(INPUT)
    for table in doc.tables:
        set_repeat_header(table)

    replace_figure_2(doc, MEMORY_CHAIN_FIG)
    replace_visible_numbering(doc)
    insert_design_figures(doc, data_packet_fig, recovery_flow_fig)
    update_existing_sections(doc)
    insert_recovery_semantics(doc)
    insert_experiment_setup_sections(doc)
    insert_results_sections(doc, recovery_results_fig, checkpoint_fig)
    update_discussion(doc)
    format_caption_paragraphs(doc)
    fix_layout_issues(doc)
    remove_trailing_empty_table_rows(doc)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
